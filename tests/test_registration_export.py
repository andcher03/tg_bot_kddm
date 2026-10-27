from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from web_admin.registration_export import (
    REGISTRATION_EXPORT_HEADERS,
    build_registrations_workbook,
)


def test_registration_export_is_plain_and_contains_registration_data():
    workbook_file = build_registrations_workbook(
        [
            {
                "user_code": "KZN-000123",
                "full_name": "Анна Иванова",
                "username": "anna_test",
                "university": "КФУ",
                "registration_date": datetime(2026, 10, 23, 12, 30),
                "status": "registered",
                "confirmation_status": "confirmed",
            }
        ],
        confirmations_enabled=True,
    )

    workbook = load_workbook(BytesIO(workbook_file.getvalue()))
    worksheet = workbook["Регистрации"]

    assert tuple(cell.value for cell in worksheet[1]) == REGISTRATION_EXPORT_HEADERS
    assert tuple(cell.value for cell in worksheet[2]) == (
        "KZN-000123",
        "Анна Иванова",
        "@anna_test",
        "КФУ",
        "23.10.2026 12:30",
        "Зарегистрирован",
        "Подтвердил участие",
    )
    assert all(
        not cell.has_style
        for row in worksheet.iter_rows()
        for cell in row
    )
    for column_cells in worksheet.columns:
        column_letter = column_cells[0].column_letter
        expected_width = max(
            len(str(cell.value or ""))
            for cell in column_cells
        ) + 2
        assert worksheet.column_dimensions[column_letter].width == expected_width


def test_registration_export_marks_confirmation_as_not_required():
    workbook_file = build_registrations_workbook(
        [
            {
                "status": "registered",
                "confirmation_status": "not_requested",
            }
        ],
        confirmations_enabled=False,
    )

    worksheet = load_workbook(BytesIO(workbook_file.getvalue()))["Регистрации"]

    assert worksheet.cell(row=2, column=7).value == "Не требуется"


def test_registration_export_escapes_excel_formulas():
    workbook_file = build_registrations_workbook(
        [
            {
                "full_name": "=HYPERLINK(\"https://example.com\")",
                "status": "registered",
            }
        ],
        confirmations_enabled=False,
    )

    worksheet = load_workbook(BytesIO(workbook_file.getvalue()))["Регистрации"]

    assert worksheet.cell(row=2, column=2).value.startswith("'=")
